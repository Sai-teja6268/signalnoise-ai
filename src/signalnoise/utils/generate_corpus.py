import os
import json
import csv
import uuid
import random
from datetime import datetime, timedelta
import pandas as pd
from docx import Document

# Define categories and templates
CATEGORIES = {
    "dependency_risk": {
        "dept": "Engineering",
        "teams": ["Payments", "Platform", "Billing", "Customer Profile", "Core API"],
        "templates": [
            "We are currently waiting for the {dep_team} team to deliver the updated {artifact}. Testing is blocked.",
            "Jira issue {jira} updated: Integration with {dep_team} API is delayed. Blocked on testing activities.",
            "Weekly Retro: Core functionality blocked by {dep_team} dependency. Deployment delayed until next sprint.",
            "Email: Customer onboarding blocked due to {dep_team} API contract mismatches. SLA is at risk.",
            "Jira ticket {jira}: Blocking dependency on {dep_team} for resource allocation. Impacting release timeline."
        ],
        "artifacts": ["API schema", "onboarding endpoints", "database connector", "auth token validation"]
    },
    "operational_outage": {
        "dept": "Operations",
        "teams": ["SRE", "Infrastructure", "DevOps", "Database Team", "Network Team"],
        "templates": [
            "Incident Report: {incident_id}. Database connection pool exhausted for {service}. 95% of requests failed.",
            "SRE Alert: High latency spike on {service}. Server CPU utilization reached 99%. Failover initiated.",
            "Post-Mortem: {service} outage lasted 45 minutes due to {root_cause}. Checkout portal was offline.",
            "System Monitor: Backup server disk space full. Nightly sync failed. Monitoring alerts delayed.",
            "SEV-1 Escalation: API servers crashed under sudden traffic load. Incident {incident_id} opened. Root cause: memory leak."
        ],
        "services": ["Checkout Gateway", "Auth Service", "Inventory DB", "Notification System"],
        "root_causes": ["memory leak", "connection pool starvation", "unoptimized SQL query", "network partition"]
    },
    "customer_escalation": {
        "dept": "Customer Success",
        "teams": ["Support", "Customer Success", "Enterprise Support", "Billing Support"],
        "templates": [
            "Support Ticket Summary: Received {count} customer complaints about {issue} this week.",
            "Account Alert: NPS score dropped from {nps_before} to {nps_after} due to system stability concerns.",
            "Enterprise Escalation: High-priority customer {customer} reported billing discrepancies. Escalating to Finance.",
            "Support retro: Ticket volume increased by {percent}% this month. Customer satisfaction has decreased.",
            "Jira Ticket {jira}: Escalation from tier-3 support. Customers unable to {action}. Refund requests rising."
        ],
        "issues": ["checkout errors", "account login failure", "invoice formatting", "missing notifications"],
        "customers": ["Acme Corp", "Globex", "Initech", "Umbrella Corp", "Stark Industries"],
        "actions": ["complete payment", "update profiles", "download invoices", "reset passwords"]
    },
    "team_burnout": {
        "dept": "Human Resources",
        "teams": ["Mobile Apps", "Frontend", "Data Pipeline", "Security Engineering"],
        "templates": [
            "Engineering Retrospective: Developers reported severe burnout from repeated weekend deployments.",
            "HR Pulse Survey: Team morale has declined by {percent}%. Heavy workload and overtime listed as main drivers.",
            "Meeting Notes: Weekend work has increased. Team leads raised attrition concerns after {resigns} key resignations.",
            "Weekly Sync: Heavy workload and high stress levels reported in {team_name} team. Capacity planning needed.",
            "Developer Retrospective: Attrition risk is high. Knowledge transfer activities have not yet been scheduled."
        ],
        "resigns": ["two senior developers", "our lead engineer", "three engineers", "the main architect"],
        "team_names": ["Frontend UI", "Analytics Data", "DevOps Platform", "Search Engine"]
    },
    "security_incident": {
        "dept": "Security",
        "teams": ["InfoSec", "Application Security", "Compliance Sec", "SecOps"],
        "templates": [
            "SecOps Alert: Firewall blocked a potential SQL injection attempt on {endpoint} endpoint.",
            "Security Incident {sec_id}: API token leaked in Github repository. Credential revocation completed.",
            "Threat Report: Phishing campaign targeted at employees. Four accounts compromised. Password reset enforced.",
            "Vulnerability Scan: High CVE found in {dep_name} dependencies. Patching scheduled for immediate release.",
            "Incident Audit {sec_id}: Unauthorized access attempt detected on database replica. IP blocked at firewall."
        ],
        "endpoints": ["/api/v1/auth/login", "/api/v1/checkout", "/api/v1/user/profile"],
        "dep_names": ["npm package", "docker base image", "pip library"]
    },
    "budget_risk": {
        "dept": "Finance",
        "teams": ["Procurement", "Cloud Infrastructure", "Finance Ops"],
        "templates": [
            "Financial Report: Vendor cost overrun for {vendor}. Cloud hosting budget exceeded by {percent}%.",
            "Operations Notice: Cloud compute costs spiked due to unoptimized {process} tasks. Action required.",
            "Budget Review: Salary inflation and contractor hiring costs exceeded Q2 allocation by ${amount}.",
            "Procurement Warning: Vendor contract renewal price increased by {percent}%. Budget adjustments needed.",
            "Finance alert: Project funding frozen due to budget constraints. Future hires delayed."
        ],
        "vendors": ["AWS", "Datadog", "Snowflake", "OpenAI"],
        "processes": ["data processing", "log indexing", "vector search indexing", "machine learning training"]
    },
    "compliance_risk": {
        "dept": "Legal & Compliance",
        "teams": ["Legal", "Compliance", "Data Governance"],
        "templates": [
            "Audit Findings: GDPR compliance audit identified gaps in {process} logs. Rectification plan needed.",
            "Compliance Notice: Privacy policy updates delayed. SOC2 compliance certification postponed.",
            "Legal Review: Unauthorized data sharing detected in test environment. Data sanitization enforced.",
            "Jira Task {jira}: Audit finding remediation. Sensitive user data found in plain text in {artifact}.",
            "Data Governance Alert: Data residency requirements violated for {region} customers. Relocation scheduled."
        ],
        "processes": ["user tracking", "payment processing", "error logging"],
        "artifacts": ["application log file", "database backups", "debug dump files"],
        "regions": ["European Union", "Asia Pacific", "Latin America"]
    },
    "quality_risk": {
        "dept": "Quality Assurance",
        "teams": ["QA", "Test Automation", "Release QA"],
        "templates": [
            "QA Report: Regression bug count increased by {percent}%. Unit test coverage has declined.",
            "Release Quality Sync: Integration test coverage dropped. High severity bug backlog reached {bug_count}.",
            "Test Alert: Automated test suite failure rate spiked to {percent}%. Main branch builds broken.",
            "Retrospective Notes: Multiple critical bugs leaked into production. QA release gate bypassed.",
            "Jira Task {jira}: Regression issues found in {feature} module. Release confidence remains low."
        ],
        "features": ["Checkout payments", "User profile update", "Search results page", "Sign-up wizard"]
    }
}

def generate_text_content(category_name: str) -> tuple[str, str, str, str]:
    config = CATEGORIES[category_name]
    dept = config["dept"]
    
    # Select random team
    team = random.choice(config["teams"]) if "teams" in config else config["dept"]
    
    # Generate content
    template = random.choice(config["templates"])
    
    # Generate template variables
    variables = {
        "dep_team": random.choice(config.get("teams", ["External Team"])),
        "artifact": random.choice(config.get("artifacts", ["resource"])),
        "jira": f"PROJ-{random.randint(100, 999)}",
        "incident_id": f"INC-{random.randint(1000, 9999)}",
        "service": random.choice(config.get("services", ["Core Service"])),
        "root_cause": random.choice(config.get("root_causes", ["unknown configuration change"])),
        "count": random.randint(10, 80),
        "issue": random.choice(config.get("issues", ["system lag"])),
        "nps_before": random.randint(35, 50),
        "nps_after": random.randint(10, 30),
        "customer": random.choice(config.get("customers", ["Client A"])),
        "percent": random.randint(15, 80),
        "action": random.choice(config.get("actions", ["perform task"])),
        "resigns": random.choice(config.get("resigns", ["a senior engineer"])),
        "team_name": random.choice(config.get("team_names", ["Core Team"])),
        "sec_id": f"SEC-{random.randint(100, 999)}",
        "endpoint": random.choice(config.get("endpoints", ["/api/v1/resource"])),
        "dep_name": random.choice(config.get("dep_names", ["module"])),
        "vendor": random.choice(config.get("vendors", ["Cloud Vendor"])),
        "process": random.choice(config.get("processes", ["job"])),
        "amount": f"{random.randint(20, 150)},000",
        "region": random.choice(config.get("regions", ["Overseas"])),
        "bug_count": random.randint(5, 30),
        "feature": random.choice(config.get("features", ["Main"]))
    }
    
    content = template.format(**variables)
    
    # Random date within last 30 days
    days_ago = random.randint(0, 30)
    dt = datetime.now() - timedelta(days=days_ago)
    date_str = dt.strftime("%Y-%m-%d")
    
    return content, team, dept, date_str

def main():
    base_dir = r"c:\Users\Admin\signalnoise-ai\data"
    os.makedirs(base_dir, exist_ok=True)
    
    print("Starting generation of Enterprise Test Corpus...")
    
    for cat_name in CATEGORIES.keys():
        cat_dir = os.path.join(base_dir, cat_name)
        os.makedirs(cat_dir, exist_ok=True)
        print(f"Generating documents for {cat_name} in {cat_dir}...")
        
        for i in range(1, 21):
            content, team, dept, date_str = generate_text_content(cat_name)
            
            # Format mixes: 10 TXT, 4 DOCX, 3 JSON, 3 CSV
            if i <= 10:
                # TXT file
                file_path = os.path.join(cat_dir, f"doc_{i}.txt")
                with open(file_path, "w", encoding="utf-8") as f:
                    # Write in a realistic format
                    f.write(f"Department: {dept}\n")
                    f.write(f"Team: {team}\n")
                    f.write(f"Date: {date_str}\n")
                    f.write(f"Status Report / Notes\n\n")
                    f.write(content + "\n")
            elif i <= 14:
                # DOCX file
                file_path = os.path.join(cat_dir, f"doc_{i}.docx")
                doc = Document()
                doc.add_heading(f"Department: {dept}", level=1)
                doc.add_heading(f"Team: {team}", level=2)
                doc.add_paragraph(f"Date: {date_str}")
                doc.add_paragraph("Enterprise Log Document Summary:")
                doc.add_paragraph(content)
                doc.save(file_path)
            elif i <= 17:
                # JSON file
                file_path = os.path.join(cat_dir, f"doc_{i}.json")
                # Add multiple items in json to verify multi-loader load
                json_data = [
                    {
                        "message": content,
                        "team": team,
                        "department": dept,
                        "date": date_str
                    },
                    {
                        "message": f"Related item: {content[:30]} verification completed.",
                        "team": team,
                        "department": dept,
                        "date": date_str
                    }
                ]
                with open(file_path, "w", encoding="utf-8") as f:
                    json.dump(json_data, f, indent=4)
            else:
                # CSV file
                file_path = os.path.join(cat_dir, f"doc_{i}.csv")
                # Add multiple rows
                csv_data = [
                    {"message": content, "team": team, "department": dept, "date": date_str},
                    {"message": f"Secondary issue logged in relation to {team} actions.", "team": team, "department": dept, "date": date_str}
                ]
                df = pd.DataFrame(csv_data)
                df.to_csv(file_path, index=False)
                
    print("Enterprise Test Corpus successfully generated!")

if __name__ == "__main__":
    main()
