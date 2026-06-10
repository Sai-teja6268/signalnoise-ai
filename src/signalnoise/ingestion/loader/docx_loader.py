import uuid
from docx import Document
from .base_loader import BaseLoader
from signalnoise.ingestion.schema import EnterpriseDocument
from .metadata_extractor import extract_metadata

class DOCXLoader(BaseLoader):
    def load(self,file_path:str):
        doc = Document(file_path)
        text = "\n".join(
            para.text for para in doc.paragraphs
        )
        
        meta = extract_metadata(text, file_path)
        
        return [
            EnterpriseDocument(
                document_id=str(uuid.uuid4()),
                source="docx",
                content=text,
                file_name=meta["file_name"],
                team=meta["team"],
                department=meta["department"],
                date=meta["date"]
            )
        ]